"""CP-07 · Generar reportes (SPVCI-03 a SPVCI-06 / PD-02 a PD-05)."""
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from apps.core.services import hoy

from .base import ConDatosDemo

D = Decimal


class CP07Reportes(ConDatosDemo):
    def setUp(self):
        self.entrar("gerente01")

    def test_cp_07_01_valor_al_corte_actual(self):
        r = self.client.get("/reportes/valor/", {"corte": hoy().isoformat()})
        self.assertEqual((r.context["res"]["total"], r.context["res"]["unidades"]), (D("87125.00"), D("427")))
        self.assertContains(r, "87,125")

    def test_cp_07_02_valor_con_corte_historico(self):
        r = self.client.get("/reportes/valor/", {"corte": "2026-08-31", "alm": self.matriz.pk})
        self.assertEqual(r.context["res"]["total"], D("69605.00"))

    def test_cp_07_03_corte_sin_datos(self):
        r = self.client.get("/reportes/valor/", {"corte": "2026-04-30"})
        self.assertEqual(r.context["f"]["errores"], [("Sin datos históricos", "No hay movimientos registrados antes del 1 may 2026.")])
        self.assertEqual(r.context["f"]["corte"], hoy())
        self.assertEqual(r.context["res"]["total"], D("87125.00"))

    def test_corte_posterior_a_hoy(self):
        r = self.client.get("/reportes/valor/", {"corte": "2099-01-01"})
        self.assertEqual(r.context["f"]["errores"][0][0], "Fecha de corte inválida")

    def test_cp_07_04_vista_previa_del_reporte(self):
        r = self.client.get("/reportes/valor/reporte/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["res"]["total"], D("87125.00"))
        self.assertContains(r, "87,125.00")
        self.assertTrue(r.context["cats"])

    def test_cp_07_05_exportar_excel(self):
        r = self.client.get("/reportes/valor/", {"formato": "xlsx", "corte": hoy().isoformat()})
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn(f'ExtinPro_ValorInventario_{hoy():%Y-%m-%d}.xlsx', r["Content-Disposition"])
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[4]], ["SKU", "Producto", "Categoría", "Unidades", "Costo unitario", "Valor"])
        total = [c.value for c in ws[ws.max_row]]
        self.assertEqual((total[3], total[5]), (427, 87125))

    def test_cp_07_06_reporte_de_ventas(self):
        r = self.client.get("/reportes/ventas/", {"desde": "2026-10-01", "hasta": "2026-10-19"})
        v = r.context["r"]
        self.assertEqual((v["ventas"], v["total"], v["iva"], v["ticket_promedio"]), (68, D("97393.02"), D("13433.52"), D("1432.25")))
        self.assertEqual(sum(m["importe"] for m in v["metodos"]), D("97393.02"))
        self.assertTrue(v["top"])
        self.assertEqual(self.client.get("/reportes/ventas/imprimir/", {"desde": "2026-10-01", "hasta": "2026-10-19"}).status_code, 200)

    def test_cp_07_07_periodo_invalido(self):
        r = self.client.get("/reportes/ventas/", {"desde": "2026-10-15", "hasta": "2026-10-01"})
        self.assertEqual(r.context["error"], "La fecha inicial no puede ser posterior a la fecha final.")
        self.assertContains(r, "La fecha inicial no puede ser posterior a la fecha final.")

    def test_panel_general(self):
        r = self.client.get("/panel/")
        self.assertEqual(r.status_code, 200)
