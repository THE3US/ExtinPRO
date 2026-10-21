
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.dateparse import parse_date

from apps.catalogo.models import CATEGORIAS, Producto
from apps.core.services import hoy
from apps.inventario.models import MovimientoInventario
from apps.usuarios.decoradores import requiere_modulo

from .services import MESES, cierre_mes_anterior, cierres_mensuales, fecha_corta, por_categoria, puntos_ventas, valuar, ventas_periodo

CERO = Decimal("0")


def inicio_historico():
    m = MovimientoInventario.objects.order_by("fecha_movimiento").values_list("fecha_movimiento", flat=True).first()
    return timezone.localtime(m).date() if m else hoy()


def _filtros_valor(request):
    dia = hoy()
    inicio = inicio_historico()
    corte = parse_date(request.GET.get("corte", "") or "") or dia
    errores = []
    if corte > dia:
        errores.append(("Fecha de corte inválida", "No puede ser posterior a hoy."))
        corte = dia
    elif corte < inicio:
        errores.append(("Sin datos históricos", f"No hay movimientos registrados antes del {inicio.day} {MESES[inicio.month - 1]} {inicio.year}."))
        corte = dia
    sucursales = request.sucursales
    alm = request.GET.get("alm", "")
    if alm.isdigit():
        sucursales = [s for s in request.sucursales if s.pk == int(alm)] or request.sucursales
        if len(sucursales) == len(request.sucursales) and len(request.sucursales) > 1:
            alm = ""
    cat = request.GET.get("cat", "")
    if cat not in CATEGORIAS:
        cat = ""
    prod = request.GET.get("prod", "")
    prod = int(prod) if prod.isdigit() else None
    return {"corte": corte, "alm": alm, "cat": cat, "prod": prod, "sucursales": sucursales, "errores": errores, "inicio": inicio}


def _datos_valor(request):
    f = _filtros_valor(request)
    dia = hoy()
    res = valuar(f["corte"], f["sucursales"], f["cat"], f["prod"], hoy=dia)
    prev_fecha = cierre_mes_anterior(f["corte"])
    prev = valuar(prev_fecha, f["sucursales"], f["cat"], f["prod"], hoy=dia) if prev_fecha >= f["inicio"] else {"total": CERO}
    delta = ((res["total"] - prev["total"]) / prev["total"] * 100) if prev["total"] else CERO
    filas = sorted([x for x in res["filas"] if x["unidades"] or f["prod"]], key=lambda x: -x["valor"])
    maximo = max((x["valor"] for x in filas), default=Decimal("1")) or Decimal("1")
    for x in filas:
        x["pct"] = (x["valor"] / res["total"] * 100) if res["total"] else CERO
        x["barra"] = (x["valor"] / maximo * 100) if maximo else CERO
    cats = por_categoria(res["filas"], res["total"])
    historico = [
        {"label": etiqueta, "value": float(valuar(d, f["sucursales"], f["cat"], f["prod"], hoy=dia)["total"]), "sub": f"Cierre {fecha_corta(d)} {d.year}"}
        for d, etiqueta in cierres_mensuales(f["corte"], f["inicio"], dia)
    ]
    nombre_alm = next((s.nombre for s in request.sucursales if str(s.pk) == f["alm"]), "Todos los almacenes")
    prod_obj = Producto.objects.filter(pk=f["prod"]).first() if f["prod"] else None
    return {
        "f": f,
        "res": res,
        "prev": prev,
        "prev_fecha": prev_fecha,
        "prev_mes": MESES[prev_fecha.month - 1],
        "delta": delta,
        "filas": filas,
        "cats": cats,
        "puntos_cat": [{"label": c["categoria"], "value": float(c["valor"]), "sub": f"{c['pct']:.1f} % del total"} for c in cats],
        "historico": historico,
        "filtro_txt": " · ".join(x for x in [nombre_alm, f["cat"] or "Todas las categorías", prod_obj.sku if prod_obj else None] if x),
        "nombre_alm": nombre_alm,
        "categorias": [c for c in CATEGORIAS if c != "Servicios"],
        "productos": Producto.objects.exclude(tipo=Producto.Tipo.SERVICIO).filter(**({"categoria": f["cat"]} if f["cat"] else {})).order_by("sku"),
        "costo_promedio": (res["total"] / res["unidades"]).quantize(Decimal("0.01")) if res["unidades"] else CERO,
        "archivo": f"ExtinPro_ValorInventario_{f['corte']:%Y-%m-%d}",
        "folio_reporte": f"RPT-INV-{f['corte']:%Y%m%d}",
    }


@requiere_modulo("valor")
def valor(request):
    """Valor de inventario (E7 · Fig. 11, 12 · SPVCI-05 / PD-04)."""
    ctx = _datos_valor(request)
    if request.GET.get("formato") == "xlsx":
        return _excel(ctx)
    return render(request, "reportes/valor.html", ctx)


@requiere_modulo("valor")
def valor_reporte(request):
    """Vista previa imprimible (Fig. 18). «Exportar PDF» = imprimir y guardar como PDF desde el navegador."""
    ctx = _datos_valor(request)
    ctx["generado"] = timezone.now()
    return render(request, "reportes/valor_reporte.html", ctx)


def _excel(ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Valor de inventario"
    ws.append([f"ExtinPro · Valor de inventario al {ctx['f']['corte']:%d/%m/%Y}"])
    ws.append([ctx["filtro_txt"] + " · valuación a precio de compra (sin IVA)"])
    ws.append([])
    ws.append(["SKU", "Producto", "Categoría", "Unidades", "Costo unitario", "Valor"])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B2F61")
    for x in ctx["filas"]:
        if x["unidades"]:
            ws.append([x["producto"].sku, x["producto"].nombre, x["producto"].categoria, float(x["unidades"]), float(x["producto"].precio_compra), float(x["valor"])])
    ws.append(["", "Total", "", float(ctx["res"]["unidades"]), "", float(ctx["res"]["total"])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=14, color="0B2F61")
    for col, ancho in zip("ABCDEF", (16, 46, 24, 12, 16, 16)):
        ws.column_dimensions[col].width = ancho
    for row in ws.iter_rows(min_row=5, min_col=5, max_col=6):
        for c in row:
            c.number_format = '"$"#,##0.00'
            c.alignment = Alignment(horizontal="right")
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{ctx["archivo"]}.xlsx"'
    return resp


# ------------------------------------------------------------------ Reporte de ventas

def _datos_ventas(request):
    dia = hoy()
    desde = parse_date(request.GET.get("desde", "") or "") if "desde" in request.GET else dia.replace(day=1)
    hasta = parse_date(request.GET.get("hasta", "") or "") if "hasta" in request.GET else dia
    error = ""
    if not desde or not hasta:
        error = "Captura el periodo completo."
    elif desde > hasta:
        error = "La fecha inicial no puede ser posterior a la fecha final."
    elif hasta > dia:
        error = "La fecha final no puede ser posterior a hoy."
    elif (hasta - desde).days > 400:
        error = "El periodo máximo del reporte es de 400 días."
    alm = request.GET.get("alm", "")
    sucursales = [s for s in request.sucursales if str(s.pk) == alm] or request.sucursales
    ctx = {"desde": desde, "hasta": hasta, "alm": alm if len(sucursales) == 1 and alm else "", "error": error,
           "nombre_alm": sucursales[0].nombre if alm and len(sucursales) == 1 else "Todos"}
    if not error:
        r = ventas_periodo(desde, hasta, sucursales)
        ctx.update({"r": r, "puntos": puntos_ventas(r["por_dia"][-31:])})
    return ctx


@requiere_modulo("reporte_ventas")
def ventas(request):
    """Reporte de ventas por periodo (E7 · Fig. 19 · SPVCI-04 / PD-03)."""
    return render(request, "reportes/ventas.html", _datos_ventas(request))


@requiere_modulo("reporte_ventas")
def ventas_imprimir(request):
    ctx = _datos_ventas(request)
    ctx["generado"] = timezone.now()
    return render(request, "reportes/ventas_reporte.html", ctx)


